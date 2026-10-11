import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from drf_spectacular.utils import extend_schema
from rest_framework import generics, permissions, status
from rest_framework.views import APIView
from rest_framework.response import Response
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.core.exceptions import ObjectDoesNotExist

from api.v1.payroll import serializers
from apps.payroll import tasks
from apps.payroll.models import PayrollPeriod, EmployeePayroll
from apps.payroll.selectors import get_attendance_data_for_period
from apps.payroll.services import finalize_employee_payroll


@extend_schema(tags=["PayrollPeriod"])
class CreatePayrollAPIView(generics.CreateAPIView):
    permission_classes = (permissions.IsAuthenticated,)
    serializer_class = serializers.PayrollPeriodSerializer

    def create(self, request, *args, **kwargs):
        try:
            enterprise = request.user.enterprise
        except ObjectDoesNotExist:
            return Response({"detail": "User is not linked to an enterprise"}, status=status.HTTP_400_BAD_REQUEST)

        if enterprise.owner_id != request.user.id:
            return Response({"detail": "Only enterprise owner can create payroll periods"}, status=status.HTTP_403_FORBIDDEN)

        return super().create(request, *args, **kwargs)


@extend_schema(tags=["PayrollPeriod"])
class PayrollPeriodListAPIView(generics.ListAPIView):
    permission_classes = (permissions.IsAuthenticated,)
    serializer_class = serializers.PayrollPeriodSerializer

    def get_queryset(self):
        return PayrollPeriod.objects.all().order_by("-end_date", "-created")


@extend_schema(tags=["PayrollPeriod"])
class RecalculatePayrollPeriodAPIView(APIView):
    permission_classes = (permissions.IsAuthenticated,)

    def post(self, request, period_id: int):
        period = get_object_or_404(PayrollPeriod, id=period_id)
        try:
            enterprise = request.user.enterprise
        except ObjectDoesNotExist:
            return Response({"detail": "User is not linked to an enterprise"}, status=status.HTTP_400_BAD_REQUEST)

        tasks.recalculate_payroll_period_task.delay(
            enterprise_id=enterprise.id,
            period_id=period.id,
        )

        return Response(
            {"detail": "Payroll recalculation started in background"},
            status=status.HTTP_202_ACCEPTED,
        )


@extend_schema(tags=["EmployeePayroll"])
class EmployeePayrollListAPIView(generics.ListAPIView):
    permission_classes = (permissions.IsAuthenticated,)
    serializer_class = serializers.EmployeePayrollSerializer

    def get_queryset(self):
        period_id = self.kwargs["period_id"]
        try:
            enterprise = self.request.user.enterprise
        except ObjectDoesNotExist:
            return EmployeePayroll.objects.none()

        return (
            EmployeePayroll.objects
            .filter(period_id=period_id, employee__enterprise=enterprise)
            .select_related("employee", "period")
            .prefetch_related("items")
            .order_by("employee__last_name", "employee__first_name")
        )


@extend_schema(tags=["EmployeePayroll"])
class EmployeePayrollDetailAPIView(generics.RetrieveAPIView):
    permission_classes = (permissions.IsAuthenticated,)
    serializer_class = serializers.EmployeePayrollSerializer

    def get_queryset(self):
        try:
            enterprise = self.request.user.enterprise
        except ObjectDoesNotExist:
            return EmployeePayroll.objects.none()
        return (
            EmployeePayroll.objects
            .filter(employee__enterprise=enterprise)
            .select_related("employee", "period")
            .prefetch_related("items")
        )


@extend_schema(tags=["EmployeePayroll"])
class FinalizeEmployeePayrollAPIView(APIView):
    permission_classes = (permissions.IsAuthenticated,)

    def post(self, request, payroll_id: int):
        try:
            enterprise = request.user.enterprise
        except ObjectDoesNotExist:
            return Response({"detail": "User is not linked to an enterprise"}, status=status.HTTP_400_BAD_REQUEST)

        payroll = get_object_or_404(
            EmployeePayroll.objects.select_related("employee"),
            id=payroll_id,
            employee__enterprise=enterprise,
        )

        finalize_employee_payroll(payroll)
        return Response({"detail": "Employee payroll finalized"}, status=status.HTTP_200_OK)


@extend_schema(
    tags=["EmployeePayroll"],
    request=serializers.FinalizePayrollListSerializer,
)
class FinalizeEmployeePayrollListAPIView(APIView):
    permission_classes = (permissions.IsAuthenticated,)

    def post(self, request):
        serializer = serializers.FinalizePayrollListSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        payroll_ids = serializer.validated_data["payroll_ids"]

        try:
            enterprise = request.user.enterprise
        except ObjectDoesNotExist:
            return Response(
                {"detail": "User is not linked to an enterprise"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        payrolls = EmployeePayroll.objects.select_related("employee").filter(
            id__in=payroll_ids,
            employee__enterprise=enterprise,
            is_finalized=False
        )

        if not payrolls.exists():
            return Response(
                {"detail": "No payrolls found"},
                status=status.HTTP_404_NOT_FOUND,
            )

        for payroll in payrolls:
            finalize_employee_payroll(payroll)

        return Response(
            {"detail": f"{payrolls.count()} payrolls finalized"},
            status=status.HTTP_200_OK,
        )


@extend_schema(tags=["PayrollPeriod"])
class PayrollAttendanceExcelAPIView(APIView):
    """
    GET /payroll-periods/<period_id>/attendance-excel/
    Returns an Excel file with daily worked hours for each employee
    within the given payroll period.

    Columns:
        1. Employee (Last First Middle)
        2..N+1. Day numbers of the period (1, 2, ..., 31) with worked hours
        Last. Total worked hours for the whole period
    """

    permission_classes = (permissions.IsAuthenticated,)

    def get(self, request, period_id: int):
        try:
            enterprise = request.user.enterprise
        except ObjectDoesNotExist:
            return Response(
                {"detail": "User is not linked to an enterprise"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        period = get_object_or_404(PayrollPeriod, id=period_id)

        employees, dates, attendance_map = get_attendance_data_for_period(
            enterprise_id=enterprise.id,
            period=period,
        )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Attendance"

        # ── Styles ──────────────────────────────────────────────
        header_font = Font(bold=True)
        center = Alignment(horizontal="center", vertical="center")
        header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        total_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

        # ── Header row ───────────────────────────────────────────
        header = ["Employee"] + [str(d.day) for d in dates] + ["Total Hours"]
        ws.append(header)

        for cell in ws[1]:
            cell.font = header_font
            cell.alignment = center
            cell.fill = header_fill

        # ── Data rows ────────────────────────────────────────────
        for employee in employees:
            full_name = f"{employee.last_name} {employee.first_name}"
            if employee.middle_name:
                full_name = f"{full_name} {employee.middle_name}"

            emp_attendance = attendance_map.get(employee.id, {})
            total_seconds = 0
            row = [full_name]

            for d in dates:
                work_seconds = emp_attendance.get(d, 0)
                total_seconds += work_seconds
                hours = round(work_seconds / 3600, 2) if work_seconds else 0
                row.append(hours)

            total_hours = round(total_seconds / 3600, 2)
            row.append(total_hours)
            ws.append(row)

            # Highlight total column cell
            total_cell = ws.cell(row=ws.max_row, column=len(header))
            total_cell.fill = total_fill
            total_cell.font = Font(bold=True)

        # ── Column widths ────────────────────────────────────────
        ws.column_dimensions["A"].width = 35
        for col_idx in range(2, len(header) + 1):
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 6

        # Last column (Total) a bit wider
        total_col_letter = openpyxl.utils.get_column_letter(len(header))
        ws.column_dimensions[total_col_letter].width = 14

        # ── Stream response ──────────────────────────────────────
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        filename = f"attendance_{period.start_date}_{period.end_date}.xlsx"
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response
